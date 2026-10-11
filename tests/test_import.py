from contextlib import closing
from datetime import datetime
import io,json,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from app import domain as d
from app.imports import HEADERS,read_registration_file
from app.web import connect
from tests import test_web as webtests


def file_bytes(rows=None):
    wb=Workbook();ws=wb.active;ws.title='Inscriptions';ws.append(HEADERS)
    for row in rows or [['Nouveau Test',datetime(2001,4,12),'F',52,'Premier cap','Oui','Non','Oui','Non','Non','Non']]:ws.append(row)
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()

class ImportTests(unittest.TestCase):
    setUp=webtests.WebTests.setUp
    tearDown=webtests.WebTests.tearDown
    token=webtests.WebTests.token
    login=webtests.WebTests.login
    post=webtests.WebTests.post
    stored=webtests.WebTests.stored

    def open_vat(self):
        c,_=self.stored();c['open']=['Dau','Vat']
        with closing(connect(self.app.config['DATABASE'])) as db:
            db.execute('UPDATE competitions SET state=? WHERE id=?',(json.dumps(c),c['id']));db.commit()

    def upload(self,data=None):
        return self.client.post('/c/'+self.c['id']+'/import',data=dict(csrf=self.token(),file=(io.BytesIO(data or file_bytes()),'inscriptions.xlsx')),content_type='multipart/form-data')

    def batch(self):
        with closing(connect(self.app.config['DATABASE'])) as db:return db.execute('SELECT id FROM import_batches').fetchone()[0]

    def confirm(self,token):
        return self.post('/c/'+self.c['id']+'/import',dict(confirm=token))

    def test_preview_then_confirm_is_atomic_pending_and_club_scoped(self):
        self.open_vat();self.login('club');r=self.upload()
        self.assertEqual(r.status_code,200,r.data);self.assertIn('Confirmer les 1 inscriptions',r.text)
        self.assertEqual(len(self.stored()[0]['people']),2)
        batch=self.batch();r=self.confirm(batch);self.assertEqual(r.status_code,200,r.data)
        c,_=self.stored();self.assertEqual(len(c['people']),3)
        self.assertEqual((c['people'][-1]['club'],c['people'][-1]['status']),('club-a','pending'))
        self.assertEqual(c['people'][-1]['disciplines'],['Dau','Vat'])
        self.assertEqual(self.confirm(batch).status_code,422)
        self.assertEqual(len(self.stored()[0]['people']),3)

    def test_bad_line_and_duplicate_prevent_all_import(self):
        self.open_vat();self.login('club')
        row=['Nouveau Test','12/04/2001','F','52,5','Premier cap','Oui','','Oui','','','']
        r=self.upload(file_bytes([row,row]));self.assertIn('Ligne 3',r.text);self.assertNotIn('name="confirm"',r.text)
        self.assertEqual(len(self.stored()[0]['people']),2)
        row[2]='Z';r=self.upload(file_bytes([row]));self.assertIn('Sexe invalide',r.text)

    def test_closed_discipline_is_rejected(self):
        self.login('club');r=self.upload();self.assertIn('fermée',r.text);self.assertNotIn('name="confirm"',r.text)

    def test_stale_preview_and_other_club_cannot_confirm(self):
        self.open_vat();self.login('club');self.upload();batch=self.batch()
        self.login('club2');self.assertEqual(self.confirm(batch).status_code,422)
        self.login('club')
        with closing(connect(self.app.config['DATABASE'])) as db:
            db.execute('UPDATE competitions SET revision=revision+1');db.commit()
        self.assertEqual(self.confirm(batch).status_code,409)
        self.assertEqual(len(self.stored()[0]['people']),2)

    def test_role_and_csrf(self):
        self.login('organizer');self.assertEqual(self.upload().status_code,403)
        self.login('club')
        self.assertEqual(self.client.post('/c/'+self.c['id']+'/import',data={}).status_code,403)

    def test_template_download_is_blank_and_examples_are_separate(self):
        self.login('club');r=self.client.get('/c/'+self.c['id']+'/import/template');self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.data));self.assertEqual(wb.sheetnames,['Inscriptions','Exemples et aide'])
        self.assertEqual([c.value for c in wb['Inscriptions'][1]],HEADERS)
        self.assertIsNone(wb['Inscriptions']['A2'].value)
        self.assertTrue(wb['Exemples et aide']['A6'].value)
        with self.assertRaises(d.RuleError):read_registration_file(r.data)
        ws=wb['Inscriptions'];example=[c.value for c in wb['Exemples et aide'][6]]
        for col,value in enumerate(example,1):ws.cell(2,col,value)
        stream=io.BytesIO();wb.save(stream)
        valid,errors=read_registration_file(stream.getvalue());self.assertEqual(len(valid),1);self.assertEqual(errors,[])

    def test_formula_bad_workbook_and_limit(self):
        row=['=1+1','12/04/2001','F',52,'Premier cap','Oui','','','','','']
        valid,errors=read_registration_file(file_bytes([row]));self.assertFalse(valid);self.assertIn('Formules',errors[0]['message'])
        with self.assertRaises(d.RuleError):read_registration_file(b'not excel')
        with self.assertRaises(d.RuleError):read_registration_file(file_bytes([row]*501))

    def test_legacy_names_normalized_without_changing_identity(self):
        c,_=self.stored();c['open']=['Combat','Lutte','Kata','Kata Arme','Combat équipe','Kata synchro']
        c['people'][0]['name']='Kata Test';c['people'][0]['disciplines']=['Combat','Kata']
        d.normalize_competition(c)
        self.assertEqual(c['open'],['Dau','Vat','Quyen','Quyen Arme','Dau équipe','Quyen synchro'])
        self.assertEqual(c['people'][0]['name'],'Kata Test')
        self.assertEqual(c['people'][0]['disciplines'],['Dau','Quyen'])
