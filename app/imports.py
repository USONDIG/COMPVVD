"""Bounded XLSX input. Values only, no formula evaluation or partial imports."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from . import domain as d

HEADERS=['Nom complet','Date de naissance','Sexe','Poids kg','Grade']+d.DISCIPLINES

def read_registration_file(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries=archive.infolist()
            d.require(len(entries)<=200 and sum(x.file_size for x in entries)<=20*1024*1024,
                      'Classeur trop volumineux une fois décompressé.')
            d.require(not any('vbaproject' in x.filename.lower() for x in entries),'Les macros ne sont pas acceptées.')
        book=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except d.RuleError:
        raise
    except Exception:
        raise d.RuleError('Fichier Excel invalide ou illisible. Utilisez le modèle .xlsx.')
    try:
        d.require('Inscriptions' in book.sheetnames,'Feuille « Inscriptions » manquante. Utilisez le modèle fourni.')
        sheet=book['Inscriptions']
        d.require((sheet.max_row or 0)<=501 and (sheet.max_column or 0)<=len(HEADERS),
                  'Maximum 500 participants et 11 colonnes. Conservez la structure du modèle.')
        # Do not trust incorrect worksheet dimensions supplied by external writers.
        sheet.reset_dimensions()
        rows=sheet.iter_rows()
        first=next(rows,None)
        d.require(first is not None and [str(c.value or '').strip() for c in first]==HEADERS,
                  'En-têtes incorrects. Conservez les colonnes du modèle, dans leur ordre.')
        valid=[];errors=[]
        for line,cells in enumerate(rows,2):
            d.require(line<=501,'Maximum 500 lignes de participants.')
            if not any(c.value is not None and str(c.value).strip() for c in cells):continue
            try:
                d.require(len(cells)<=len(HEADERS),'Colonnes supplémentaires non autorisées.')
                d.require(not any(c.data_type in ('f','e') for c in cells),'Formules et erreurs Excel interdites : utilisez des valeurs.')
                values=[c.value for c in cells]+[None]*(len(HEADERS)-len(cells))
                d.require(all(len(str(v))<=200 for v in values if v is not None),'Une cellule dépasse 200 caractères.')
                name,birth,sex,weight,grade=values[:5]
                if isinstance(birth,(datetime,date)):birth=birth.date().isoformat() if isinstance(birth,datetime) else birth.isoformat()
                else:
                    text=str(birth or '').strip()
                    try:birth=datetime.strptime(text,'%d/%m/%Y').date().isoformat() if '/' in text else date.fromisoformat(text).isoformat()
                    except ValueError:raise d.RuleError('Date de naissance attendue au format JJ/MM/AAAA.')
                grade=str(grade or '').strip()
                d.require(grade in d.GRADES,'Grade inconnu : choisissez un grade du modèle.')
                selected=[]
                for disc,value in zip(d.DISCIPLINES,values[5:]):
                    flag=str(value or '').strip().lower()
                    d.require(flag in ('','oui','non','x','1','0'),'Disciplines : indiquez Oui, Non ou laissez vide.')
                    if flag in ('oui','x','1'):selected.append(disc)
                fields=[('name',str(name or '').strip()),('birth',birth),('sex',str(sex or '').strip().upper()),
                        ('weight',str(weight or '').strip().replace(',','.')),('grade',str(d.GRADES.index(grade)))]+[('disciplines',x) for x in selected]
                valid.append((line,fields))
            except d.RuleError as exc:errors.append(dict(line=line,message=str(exc)))
        d.require(valid or errors,'La feuille Inscriptions est vide. Les exemples ne sont pas importés.')
        return valid,errors
    except d.RuleError:raise
    except Exception:raise d.RuleError('Le contenu du classeur est illisible. Enregistrez-le à nouveau en .xlsx.')
    finally:book.close()
